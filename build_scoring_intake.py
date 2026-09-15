# -*- coding: utf-8 -*-
"""
Build the scoring intake workbook for the active HEM protocols.

WHAT THIS DOES AND DOES NOT DO
  The rubric has 37 items across 8 domains. The portfolio/OnCore/CT.gov data
  can defensibly fill only a handful of them -- all structural (Domain 1/4).
  Domains 3, 5, 6, 7 and most of 8 -- visit burden, specimen handling, safety
  reporting, qualitative load, abstraction volume -- are not in any data source
  here and cannot be inferred from a title.

  So this builds an INTAKE SHEET, not scores. Derivable items are pre-filled
  with the rule and source recorded for each; everything else is left blank.
  Blank means blank: batch_score.js refuses to total a protocol until all 37
  are present, because the tool's own loader treats an absent item as 0 and
  would silently understate the score.

  Part B (WU/FTE) additionally needs participant counts broken out by status
  row. The portfolio carries Total Accrual only, so Part B is not computable
  for any protocol here without a second data pull.

USAGE  python3 build_scoring_intake.py [out_dir]
"""
import os, sys, re, datetime, json
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT_DIR = sys.argv[1] if len(sys.argv) > 1 else (
    "/Users/jhodges/Library/CloudStorage/"
    "OneDrive-St.JudeChildren'sResearchHospital/HEM Protocol Reviews")
TODAY = datetime.date.today().isoformat()
BASE = "/Users/jhodges/Library/Mobile Documents/com~apple~CloudDocs/Documents"
PORT = BASE + "/HEM_Clinical_Trials_Dashboard/CTM Protocol Dashboards/HEM_Protocol_Portfolio.xlsx"
MSTR = BASE + "/HEM_Clinical_Trials_Dashboard/HEMProtocolReview/HemProtocolReviewMASTER.xlsx"
REPO = BASE + "/GitHub/hem-protocol-scoring-tool"

# ── rubric items, dumped from the real app.js by dump_rubric.js ──────────
# Parsing JS with a regex missed two items whose notes are double-quoted, so
# the tables come from the engine itself. Regenerate with:
#   jsc dump_rubric.js -- <repo> rubric.json
RUBRIC = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "rubric.json")))
DOMAIN_TITLES = [(d["id"], d["title"]) for d in RUBRIC["domains"]]
ITEMS = [(d["id"], it["id"], it["label"], it["max"], it["note"])
         for d in RUBRIC["domains"] for it in d["items"]]
assert len(ITEMS) == 37, "expected 37 items, got %d" % len(ITEMS)
PART_A_MAX = RUBRIC["partAMax"]

# ── data ──────────────────────────────────────────────────────────────────
pf = openpyxl.load_workbook(PORT, data_only=True)["Portfolio"]
PH = [c.value for c in pf[1]]
prot = [dict(zip(PH, r)) for r in pf.iter_rows(min_row=2, values_only=True) if r[0]]
active = [r for r in prot if str(r["Status Group"]) != "IRB Closure"]

mw = openpyxl.load_workbook(MSTR, data_only=True)
oc = mw["OnCoreProtSearch"]; OH = [str(c.value) for c in oc[1]]
onc = {}
for r in oc.iter_rows(min_row=2, values_only=True):
    d = dict(zip(OH, r))
    if d.get("Protocol No."):
        onc[str(d["Protocol No."]).strip().upper()] = d
        for alt in str(d.get("Additional Protocol Numbers") or "").split(";"):
            if alt.strip(): onc.setdefault(alt.strip().upper(), d)

def S(v): return str(v).strip() if v not in (None, "") else ""

# ── derivation rules: each returns (value, rule_text) or None ─────────────
def d_sponsor_type(p, o):
    st = S(p.get("Sponsor Type")) or S(o.get("Sponsor Type") if o else "")
    if st == "Institutional": return 0, "Sponsor Type = Institutional -> investigator-initiated/internal (anchor 0)"
    if st == "National":      return 2, "Sponsor Type = National -> NIH/federal (mid anchor; 2 of 0-3)"
    if st == "Industry":      return 3, "Sponsor Type = Industry (max anchor)"
    return None

def d_site_role(p, o):
    if not o: return None
    ms = S(o.get("Multi-site Trial"))
    if ms == "N": return 0, "OnCore Multi-site Trial = N -> single-site (anchor 0)"
    if ms == "Y":
        iit = S(o.get("Investigator Initiated Protocol"))
        if iit == "Y": return 5, "OnCore Multi-site = Y and Investigator Initiated = Y -> SJ likely coordinating centre (max anchor). VERIFY."
        return 3, "OnCore Multi-site = Y, not investigator-initiated -> participating site (mid anchor; 3 of 0-5)"
    return None

def d_reg_status(p, o):
    if not o: return None
    ind = S(o.get("IND Ids")); drug = S(o.get("Investigational Drug"))
    iit = S(o.get("Investigator Initiated Protocol"))
    if ind and ind.lower() != "none":
        if iit == "Y": return 5, "OnCore IND Ids present and investigator-initiated -> active IND held by St. Jude (max). VERIFY holder."
        return 3, "OnCore IND Ids present, not investigator-initiated -> regulated, IND held externally (mid; 3 of 0-5). VERIFY."
    if drug == "No": return 0, "OnCore Investigational Drug = No -> not FDA-regulated (anchor 0). VERIFY no device/IDE."
    return None

def d_ip_handling(p, o):
    gt = S(p.get("Gene Therapy"))
    drug = S(o.get("Investigational Drug")) if o else ""
    if gt == "Y": return 4, "Portfolio Gene Therapy = Y -> cell/gene product with chain-of-custody (max anchor)"
    if drug == "Yes": return 2, "OnCore Investigational Drug = Yes, not gene therapy -> conventional IP (mid; 2 of 0-4)"
    if drug == "No" and S(p.get("Therapeutic Group")) == "Non-Therapeutic":
        return 0, "OnCore Investigational Drug = No and Therapeutic Group = Non-Therapeutic -> no IP (anchor 0)"
    return None

def d_blinding_rand(p, o):
    t = (S(p.get("Title")) + " " + S(p.get("Short Title"))).lower()
    if not t: return None
    if "double-blind" in t or "double blind" in t:
        return 3, "Title states double-blind (max anchor)"
    if "open-label" in t or "open label" in t:
        return 0, "Title states open-label (anchor 0)"
    if "randomi" in t:
        return 2, "Title states randomised but not double-blind (2 of 0-3). VERIFY masking."
    if S(p.get("Therapeutic Group")) == "Non-Therapeutic":
        return 0, "Non-therapeutic study -> no blinding/randomisation (anchor 0)"
    return None

def d_pediatric(p, o):
    if not o: return None
    ag = S(o.get("Age Group"))
    if ag == "Adults": return 0, "OnCore Age Group = Adults -> adults only (anchor 0)"
    return None   # Children/Both cannot tell assent-tier structure

RULES = [("sponsor_type", d_sponsor_type), ("site_role", d_site_role),
         ("reg_status", d_reg_status), ("ip_handling", d_ip_handling),
         ("blinding_rand", d_blinding_rand), ("pediatric_assent", d_pediatric)]

# ── triage (categorical, deliberately not a score) ────────────────────────
def triage(p, o):
    sig = []
    if S(p.get("Gene Therapy")) == "Y": sig.append("gene therapy")
    if S(p.get("Sponsor Type")) == "Industry": sig.append("industry-sponsored")
    if o and S(o.get("Multi-site Trial")) == "Y": sig.append("multi-site")
    if o and S(o.get("IND Ids")) and S(o.get("IND Ids")).lower() != "none": sig.append("IND")
    if S(p.get("Phase")) in ("I", "I/II", "II", "II/III", "III", "IV"): sig.append("phase " + S(p.get("Phase")))
    if S(p.get("Therapeutic Group")) == "Therapeutic": sig.append("therapeutic")
    n = len(sig)
    band = "Score first" if n >= 4 else ("Score second" if n >= 2 else "Low structural signal")
    return band, ", ".join(sig) if sig else "none of the structural markers"

# ── build ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
HDR = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1F4E79")
PREFILL = PatternFill("solid", fgColor="E8F5E9"); BLANKF = PatternFill("solid", fgColor="FFF8E1")
WRAP = Alignment(wrap_text=True, vertical="top")

ws = wb.active; ws.title = "Intake"
meta_cols = ["Protocol", "PI", "Status", "Disease Group", "Therapeutic Group", "Phase",
             "Sponsor Type", "Gene Therapy", "OnCore record?", "Items pre-filled",
             "Items still needed", "Triage"]
ws.append(meta_cols + ["%s | %s (0-%d)" % (d, i, mx) for d, i, lbl, mx, nt in ITEMS])

log = []
counts = Counter()
for p in sorted(active, key=lambda x: str(x["Protocol"])):
    pid = str(p["Protocol"]).strip()
    o = onc.get(pid.upper())
    vals = {}
    for iid, fn in RULES:
        got = fn(p, o)
        if got:
            vals[iid] = got[0]
            log.append([pid, iid, got[0], got[1]])
    band, why = triage(p, o)
    counts[band] += 1
    row = [pid, S(p.get("PI")), S(p.get("Status Group")), S(p.get("Disease Group")),
           S(p.get("Therapeutic Group")), S(p.get("Phase")), S(p.get("Sponsor Type")),
           S(p.get("Gene Therapy")), "yes" if o else "no",
           len(vals), 37 - len(vals), band]
    row += [vals.get(i, "") for d, i, lbl, mx, nt in ITEMS]
    ws.append(row)

for i, (d, iid, lbl, mx, nt) in enumerate(ITEMS):
    col = len(meta_cols) + 1 + i
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=col).fill = PREFILL if ws.cell(row=r, column=col).value != "" else BLANKF

ws2 = wb.create_sheet("Item reference")
ws2.append(["#", "Domain", "Item id", "Item", "Max", "Anchors", "Pre-filled by this build?"])
prefillable = {i for i, _ in RULES}
for n, (d, iid, lbl, mx, nt) in enumerate(ITEMS, 1):
    ws2.append([n, dict(DOMAIN_TITLES).get(d, d), iid, lbl, mx, nt,
                "yes - verify" if iid in prefillable else "NO - must be scored from the protocol"])

ws3 = wb.create_sheet("Derivation log")
ws3.append(["Protocol", "Item id", "Value", "Rule and source"])
for r in log: ws3.append(r)

ws4 = wb.create_sheet("Triage")
ws4.append(["Protocol", "PI", "Disease Group", "Status", "Band", "Structural markers present"])
for p in sorted(active, key=lambda x: str(x["Protocol"])):
    o = onc.get(str(p["Protocol"]).strip().upper())
    band, why = triage(p, o)
    ws4.append([str(p["Protocol"]).strip(), S(p.get("PI")), S(p.get("Disease Group")),
                S(p.get("Status Group")), band, why])

for sh, widths in (("Intake", [16,22,18,26,16,8,14,8,13,13,15,18]),
                   ("Item reference", [4,44,20,44,6,86,34]),
                   ("Derivation log", [16,18,7,96]),
                   ("Triage", [18,24,28,18,20,64])):
    s = wb[sh]
    for i, w in enumerate(widths, 1): s.column_dimensions[get_column_letter(i)].width = w
    if sh == "Intake":
        for i in range(len(widths)+1, len(widths)+38): s.column_dimensions[get_column_letter(i)].width = 13
    for c in s[1]: c.font = HDR; c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical="center")
    s.freeze_panes = "A2"; s.auto_filter.ref = s.dimensions
    for row in s.iter_rows(min_row=2):
        for c in row: c.alignment = WRAP

out = os.path.join(OUT_DIR, "Protocol_Scoring_Intake_%s.xlsx" % TODAY)
wb.save(out)
print("WROTE:", out)
print("active protocols (not IRB Closure): %d | with OnCore record: %d"
      % (len(active), sum(1 for p in active if str(p["Protocol"]).strip().upper() in onc)))
print("pre-filled cells: %d of %d (%.1f%%)" % (len(log), len(active)*37, 100.0*len(log)/(len(active)*37)))
print("triage:", dict(counts))
