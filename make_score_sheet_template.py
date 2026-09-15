# -*- coding: utf-8 -*-
"""Generate the corrected Hematology protocol score sheet template."""
import docx, os
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = ("/Users/jhodges/Library/CloudStorage/"
        "OneDrive-St.JudeChildren'sResearchHospital/HEM Protocol Reviews")
NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)

d = docx.Document()
for s in d.sections:
    s.top_margin = s.bottom_margin = Inches(0.6)
    s.left_margin = s.right_margin = Inches(0.7)
st = d.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(10)

def fixed(table, widths):
    """Force fixed layout and apply column widths (inches).

    LibreOffice and Word lay out from w:tblGrid, which python-docx does not
    update when you set cell.width -- so rewrite the grid as well as the cells.
    """
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    for el in tblPr.findall(qn("w:tblLayout")): tblPr.remove(el)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    # rewrite w:tblGrid
    for grid in tbl.findall(qn("w:tblGrid")): tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(int(w * 1440))); grid.append(gc)
    tblPr.addnext(grid)
    for row in table.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells): row.cells[i].width = Inches(w)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd"); el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor); tcPr.append(el)

def para(text="", size=10, bold=False, color=None, after=4, align=None, italic=False):
    p = d.add_paragraph(); r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    return p

# ---------------- Title ----------------
para("HEMATOLOGY PROTOCOL SCORE SHEET", 15, True, NAVY, after=1)
para("Department of Hematology · Protocol Review Committee", 9, False, GREY, after=8)

# ---------------- Metadata ----------------
meta = d.add_table(rows=4, cols=4); meta.style = "Table Grid"
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
fields = [("Protocol", ""), ("PI", ""),
          ("Sponsor", ""), ("Funding", ""),
          ("Reviewer", ""), ("Review date", ""),
          ("Protocol version / date", ""), ("Study type", "")]
for i, (label, val) in enumerate(fields):
    r, c = divmod(i, 2)
    lc, vc = meta.cell(r, c*2), meta.cell(r, c*2+1)
    lc.text = label; vc.text = val
    shade(lc, "EDF2F7")
    for p in lc.paragraphs:
        for run in p.runs: run.bold = True; run.font.size = Pt(9)
    for p in vc.paragraphs:
        for run in p.runs: run.font.size = Pt(9)
fixed(meta, [1.5, 2.25, 1.4, 1.95])

d.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------------- Scoring instruction box ----------------
box = d.add_table(rows=1, cols=1); box.style = "Table Grid"
bc = box.cell(0, 0); shade(bc, "FFF8E1")
bc.text = ""
p = bc.paragraphs[0]
r = p.add_run("How to score  ")
r.bold = True; r.font.size = Pt(10); r.font.color.rgb = NAVY
r = p.add_run("Grade each criterion 1–5 (1 = lowest, 5 = highest quality). "
              "Maximum possible score is 6 × 5 = ")
r.font.size = Pt(9)
r = p.add_run("30"); r.bold = True; r.font.size = Pt(9)
r = p.add_run("."); r.font.size = Pt(9)
p2 = bc.add_paragraph()
r = p2.add_run("Marking a criterion NA  "); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = NAVY
r = p2.add_run("If a criterion genuinely does not apply, write ")
r.font.size = Pt(9)
r = p2.add_run("NA"); r.bold = True; r.font.size = Pt(9)
r = p2.add_run(" — do not score it 0 and do not score it 3. An NA criterion is "
               "removed from the denominator entirely. Score the study on what applied.")
r.font.size = Pt(9)
p3 = bc.add_paragraph()
r = p3.add_run("Worked example  "); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = NAVY
r = p3.add_run("A sponsor trial where Statistical plan and Methods are both NA, scoring "
               "5 · 5 · 5 · 3 on the rest → total ")
r.font.size = Pt(9)
r = p3.add_run("18 / 20 (90%)"); r.bold = True; r.font.size = Pt(9)
r = p3.add_run(", because 4 criteria × 5 = 20. Always record the percentage — it is the "
               "only figure comparable across protocols.")
r.font.size = Pt(9)
for p in bc.paragraphs: p.paragraph_format.space_after = Pt(3)

d.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------------- Scoring table ----------------
CRIT = [
    ("1. Statistical plan", "Is the analysis plan adequate and powered for the question asked?"),
    ("2. Methods",          "Are the design, endpoints, and data collection sound and clearly described?"),
    ("3. Mission",          "Does this advance St. Jude's mission in non-malignant hematology?"),
    ("4. Innovation",       "Is this novel, or does it duplicate work already done here or elsewhere?"),
    ("5. Relevance",        "Does it address a real unmet need for our patient population?"),
    ("6. Feasibility",      "Can we actually accrue, fund, and operationally deliver this?"),
]
tbl = d.add_table(rows=1 + len(CRIT), cols=3); tbl.style = "Table Grid"
hdr = ["Criterion", "Score\n(1–5 or NA)", "Comment / justification"]
for i, h in enumerate(hdr):
    c = tbl.cell(0, i); c.text = h; shade(c, "1F4E79")
    for p in c.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True; run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
for i, (name, prompt) in enumerate(CRIT, 1):
    c0 = tbl.cell(i, 0); c0.text = ""
    p = c0.paragraphs[0]; r = p.add_run(name); r.bold = True; r.font.size = Pt(9.5)
    p2 = c0.add_paragraph(); r = p2.add_run(prompt)
    r.font.size = Pt(8); r.italic = True; r.font.color.rgb = GREY
    p2.paragraph_format.space_after = Pt(2)
    tbl.cell(i, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
fixed(tbl, [2.45, 0.85, 3.80])
for row in tbl.rows:
    if row is not tbl.rows[0]: row.height = Inches(0.42)

# ---------------- Total ----------------
tot = d.add_table(rows=1, cols=4); tot.style = "Table Grid"
labels = ["Sum of scored criteria", "Number of criteria scored (not NA)",
          "Denominator (n × 5)", "TOTAL   ___ / ___  =  ____ %"]
for i, l in enumerate(labels):
    c = tot.cell(0, i); c.text = ""
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(l); r.font.size = Pt(8.5); r.bold = (i == 3)
    if i == 3: r.font.color.rgb = NAVY
    shade(c, "EDF2F7" if i < 3 else "D6E4F0")
    c.add_paragraph()
fixed(tot, [1.7, 2.0, 1.5, 1.9])

d.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------------- Decision block ----------------
para("COMMITTEE DECISION", 11, True, NAVY, after=2)
para("Required. This is what is recorded in CumulativeReviewtracking.xlsx — a score sheet "
     "without a decision leaves no usable record of the review.", 8, False, GREY, after=4, italic=True)

dec = d.add_table(rows=5, cols=2); dec.style = "Table Grid"
rows = [
 ("Recommendation",
  "☐ Open — high priority     ☐ Open — standard     ☐ Open with conditions\n"
  "☐ Tabled — revisit on ______________     ☐ Decline     ☐ More information needed"),
 ("Conditions / actions required",
  "What must happen before this opens, or before it is revisited?"),
 ("Owner", "Who is responsible for the follow-up action?"),
 ("Due by", ""),
 ("Resulting comment for the tracker",
  "One or two sentences. This text goes straight into the tracker."),
]
for i, (label, hint) in enumerate(rows):
    lc, vc = dec.cell(i, 0), dec.cell(i, 1)
    lc.text = ""; p = lc.paragraphs[0]
    r = p.add_run(label); r.bold = True; r.font.size = Pt(9)
    shade(lc, "EDF2F7")
    vc.text = ""
    if hint and i in (0,):
        p = vc.paragraphs[0]; r = p.add_run(hint); r.font.size = Pt(9)
    elif hint:
        p = vc.paragraphs[0]; r = p.add_run(hint)
        r.font.size = Pt(8); r.italic = True; r.font.color.rgb = GREY
    vc.add_paragraph()
fixed(dec, [2.0, 5.1])
dec.rows[1].height = Inches(0.55); dec.rows[4].height = Inches(0.55)

# ================= PAGE 2: anchors =================
d.add_section(WD_SECTION.NEW_PAGE)
para("SCORING ANCHORS", 14, True, NAVY, after=1)
para("Adapted from the department's Protocol Assessment Workbook rubric so that the six "
     "criteria used in review map onto the standards already agreed there. Use these to keep "
     "scoring consistent between reviewers.", 8.5, False, GREY, after=8, italic=True)

ANCH = [
 ("1. Statistical plan",
  ["No statistical plan", "Basic descriptive only", "Appropriate analysis specified",
   "Pre-specified SAP", "Detailed SAP with interim analysis plan"],
  "Score NA only for expanded-access, registry intake, or single-patient protocols where no "
  "analysis is proposed. A missing power calculation is a score of 1, not NA."),
 ("2. Methods",
  ["Ad hoc collection; endpoints unclear", "Standard forms, no validation; primary poorly defined",
   "Validated instruments; clear primary", "EDC with edit checks; well-defined hierarchy",
   "Validated EDC, real-time monitoring; pre-specified hierarchy"],
  "Covers design, endpoints, inclusion/exclusion, and data collection together."),
 ("3. Mission",
  ["Not applicable to our populations", "Tangentially related", "Indirectly applicable",
   "Directly applicable to a subset", "Central to non-malignant hematology care at St. Jude"],
  "Ask whether St. Jude is the right place for this study, not whether the science is good."),
 ("4. Innovation",
  ["Duplicates existing work", "Incremental contribution", "Moderate innovation",
   "Significant innovation", "Transformative approach"],
  "Duplication of an existing internal study is a 1 — flag the conflicting protocol by name."),
 ("5. Relevance",
  ["No unmet need", "Minor clinical gap", "Moderate unmet need",
   "Significant unmet need", "Critical / urgent unmet need; no effective treatments"],
  "Relevance is about the patients' need, not the study's quality."),
 ("6. Feasibility",
  ["<5 eligible/yr, unfunded, major operational lift", "5–10 eligible/yr, minimal funding",
   "10–25 eligible/yr, partial funding", "25–50 eligible/yr, substantially funded",
   ">50 eligible/yr, fully funded, existing capacity"],
  "State the actual eligible-patient estimate and the funding source in your comment. "
  "Feasibility is the criterion that most often decides the outcome."),
]
at = d.add_table(rows=1 + len(ANCH), cols=6); at.style = "Table Grid"
heads = ["Criterion", "1", "2", "3", "4", "5"]
for i, h in enumerate(heads):
    c = at.cell(0, i); c.text = h; shade(c, "1F4E79")
    for p in c.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True; run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
for i, (name, anchors, note) in enumerate(ANCH, 1):
    c = at.cell(i, 0); c.text = ""
    p = c.paragraphs[0]; r = p.add_run(name); r.bold = True; r.font.size = Pt(8.5)
    p.paragraph_format.space_after = Pt(1)
    for j, a in enumerate(anchors, 1):
        cc = at.cell(i, j); cc.text = ""
        p = cc.paragraphs[0]; r = p.add_run(a); r.font.size = Pt(7.5)
        p.paragraph_format.space_after = Pt(1)
fixed(at, [1.55, 1.12, 1.12, 1.12, 1.12, 1.12])

d.add_paragraph().paragraph_format.space_after = Pt(4)
para("Notes on individual criteria", 10, True, NAVY, after=3)
for name, anchors, note in ANCH:
    p = d.add_paragraph()
    r = p.add_run(name.split(". ", 1)[1] + "  "); r.bold = True; r.font.size = Pt(8.5)
    r = p.add_run(note); r.font.size = Pt(8.5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.15)

d.add_paragraph().paragraph_format.space_after = Pt(6)
para("Why the denominator changed", 11, True, NAVY, after=2)
para("The previous sheet's header read \"possible scores 5-25\" while listing six criteria, "
     "and reviewers resolved the contradiction three different ways — dropping NA from the "
     "denominator, scoring NA as 0, and scoring NA as 3. Across the 14 protocols reviewed "
     "between 20 Aug and 1 Oct 2025 this produced totals out of 15, 20, 25 and 30, none of "
     "them comparable. Dropping NA from the denominator is now the single rule; recording the "
     "percentage makes a sponsor trial scored out of 20 comparable with an investigator study "
     "scored out of 30.", 9, after=4)

out = os.path.join(ROOT, "Protocol Score Sheet TEMPLATE v2.docx")
d.save(out)
print("WROTE:", out, os.path.getsize(out), "bytes")
